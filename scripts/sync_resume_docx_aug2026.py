"""
sync_resume_docx_aug2026.py

Brings Resume.v3.7.docx in line with the current resume_site/templates/resume.html
content as of the Aug 2026 wording/date/label pass (EMRTS end date, full month names,
"Lowe's" naming, "Contract position." labels, and several Tech lines that had drifted
out of sync or were missing/duplicated).

Run from the project root:

    conda run -n flask_env python scripts/sync_resume_docx_aug2026.py

A backup is written to static/files/Resume.v3.7.backup2.docx before any changes.
"""

import shutil
from pathlib import Path
from docx import Document

DOCX_PATH = Path("static/files/Resume.v3.7.docx")
BACKUP_PATH = Path("static/files/Resume.v3.7.backup2.docx")


def replace_para_text(p, new_text: str) -> None:
    runs = p.runs
    if not runs:
        p.add_run(new_text)
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""


def append_tech(p, tech_value: str) -> None:
    """Append ' Tech: <value>' to a paragraph as a new bold-label run pair."""
    p.add_run(" ")
    tech_run = p.add_run("Tech:")
    tech_run.bold = True
    p.add_run(" " + tech_value)


def main():
    if not DOCX_PATH.exists():
        raise FileNotFoundError(f"Not found: {DOCX_PATH}")

    shutil.copy2(DOCX_PATH, BACKUP_PATH)
    print(f"Backup written to {BACKUP_PATH}")

    doc = Document(str(DOCX_PATH))
    P = doc.paragraphs

    # --- Summary ---
    replace_para_text(P[6], (
        "Thirty years of professional software development spanning telecommunications, "
        "healthcare IT, and data systems. Primary focus has been backend engineering on "
        "Unix/Linux platforms, with experience across a broad range of languages, databases, "
        "and system architectures. More recent work includes web development, functional "
        "programming, and AI-assisted development and prompt engineering."
    ))
    replace_para_text(P[7], (
        "Proof-of-concept development — particularly in healthcare IT — has been a "
        "recurring part of the work, as has intern mentoring and occasional collaboration with "
        "academic and research teams. Interested in expanding on that going forward."
    ))

    # --- Technical Skills: drop stray trailing punctuation ---
    P[18].runs[1].text = " Django, Flask, Phoenix, SQLAlchemy"
    P[21].runs[1].text = " Git, Docker, Node.js, Shell Scripting, RESTful APIs"

    # --- EMRTS block ---
    replace_para_text(P[28], (
        "Software engineer working with a healthcare IT startup across six engagements, "
        "developing open-source tools for medical records, billing, provider enrollment, and "
        "claims. Work has centered on proof-of-concept development and R&D, with occasional "
        "collaboration with academic and professional partners. Contracts listed below fill the "
        "gaps between EMRTS engagements."
    ))
    replace_para_text(P[29], (
        "July 2025–August 2026: Developed a proof-of-concept Provider Vault medical "
        "enrollment website, combining Elixir/Phoenix for scalability and concurrency with "
        "Python for AI-assisted development and prompt engineering. Subsequently developed the "
        "Medicaid Claims Checker, a fraud detection tool written in Haskell (core logic) with "
        "Elixir/Phoenix for I/O, submitted via RFI (Request for Information) to several state "
        "governments. Mentoring interns on complementary healthcare IT projects, with topics "
        "including machine learning and prompt engineering. Tech: Haskell, Elixir, Phoenix, "
        "Python 3.11, PostgreSQL"
    ))
    replace_para_text(P[30], (
        "August 2024–December 2024: Collaborated with interns on a Medicaid prediction "
        "research project, introducing foundational data science concepts and workflows. Led a "
        "follow-up initiative exploring cryptocurrency, developing a Bitcoin indexer prototype. "
        "Tech: Python 3.x, Flask, PostgreSQL"
    ))
    replace_para_text(P[31], (
        "Interim contracts: CSpire, Jackson MS (July 2018–January 2019); Orchid Orthopedic "
        "Solutions, Remote (February 2020–December 2020); Quantuvos, Remote (July "
        "2021–February 2022) — see entries below."
    ))
    replace_para_text(P[34], (
        "Interim contracts: Hilton, Memphis TN (April 2012–August 2012); Lowe's, "
        "Mooresville NC (August 2012–June 2013) — see entries below."
    ))
    replace_para_text(P[36], (
        "Interim contract: Hewlett-Packard, Nashville TN (May 2010–December 2011) — "
        "see entry below."
    ))

    # --- Quantuvos ---
    P[41].runs[0].text = "Remote | July 2021 – February 2022"
    replace_para_text(P[42], (
        "Technical consulting for Quantuvos, a coaching-platform startup. Reviewed and "
        "evaluated code for the client-facing MVP, supporting coaches and clients in scheduling "
        "workshops and sessions, identifying issues and recommending improvements. Contract "
        "position."
    ))
    append_tech(P[42], "Python 3.9, Django 3.1, PostgreSQL 14.1, Docker")

    # --- Orchid ---
    P[45].runs[0].text = "Remote | February 2020 – December 2020"
    replace_para_text(P[46], (
        "Contract work preparing Word, Excel, and PDF documents for archival submission to EtQ "
        "Reliance. Programmatically inserted tags and formatting into them and validated "
        "spreadsheet metadata for accuracy and completeness. Contract position."
    ))
    append_tech(P[46], "Python 3.8, C#, Windows/macOS")

    # --- CSpire ---
    P[49].runs[1].text = "ackson, MS | July 2018 – January 2019"
    append_tech(P[50], "Python 2.7, Flask, SQLAlchemy, MySQL, PostgreSQL")

    # --- Lowe's (2012, EMV) ---
    P[51].runs[0].text = "Lowe’s"
    P[51].runs[1].text = ""
    P[52].runs[0].text = "Mooresville, NC | August 2012 – June 2013"
    replace_para_text(P[53], (
        "Refactored and extended legacy C code to support EMV (chip-based credit and debit "
        "card) transaction processing as part of Lowe’s transition away from magnetic "
        "stripe (swipe) cards. Work included code analysis, bug fixes, and documentation. "
        "Contract position."
    ))

    # --- Hilton ---
    P[56].runs[0].text = "Memphis, TN | April 2012 – August 2012"
    P[57].runs[0].text = (
        "Provided production support for Hilton’s Central Reservation System, contributing "
        "bug fixes and maintenance using C, Perl, and shell scripting. Contract position."
    )

    # --- Hewlett-Packard ---
    P[60].runs[0].text = "Nashville, TN | May 2010 – December 2011"

    # --- ACS ---
    P[68].runs[0].text = "Nashville, TN | February 2009 – June 2009"
    replace_para_text(P[69], (
        "Migrated MMIS (Medicaid Management Information System) software for TennCare from "
        "Hewlett-Packard’s infrastructure to ACS (Affiliated Computer Services). Developed "
        "automation scripts and analyzed legacy C modules. Contract position."
    ))

    # --- Nationwide (dates/desc/Tech merged into one paragraph in this docx) ---
    P[72].runs[0].text = "Columbus, OH | October 2008 – February 2009"
    P[72].runs[2].text = (
        "Analyzed annuity and financial data daily to detect anomalies in high-volume "
        "datasets. Proposed and developed a Python-based data validation and reporting tool to "
        "modernize and replace aging shell scripts used in production. Contract position."
    )
    P[72].runs[5].text = " Solaris, Oracle, proprietary DSL"

    # --- Lexmark ---
    P[75].runs[0].text = "Lexington, KY | May 2007 – October 2008"
    append_tech(P[76], "C, Python 2.6, RHEL")

    # --- IBM Internet Security Systems (had Lexmark's text duplicated by mistake) ---
    P[79].runs[0].text = "Atlanta, GA | January 2007 – May 2007"
    replace_para_text(P[80], (
        "Developed tools for porting and rebuilding C/C++ applications from Microsoft Visual "
        "Studio to Unix environments, using Visual Studio build artifacts rather than "
        "traditional makefile-based builds. Contract position."
    ))
    append_tech(P[80], "C/C++, Visual Studio, macOS, Ubuntu Linux, Windows XP")

    # --- Lowe's (2006, Canadian store) ---
    P[82].runs[0].text = "Lowe’s"
    P[83].runs[0].text = "Wilkesboro, NC | February 2006 – December 2006"
    P[84].runs[0].text = ""
    P[84].runs[1].text = (
        "Adapted Lowe’s point-of-sale software for the company’s first Canadian "
        "store, modifying C code for local currency, card processing, and regulatory "
        "compliance. Contract position."
    )

    # --- Cingular Wireless ---
    P[88].runs[0].text = "Alpharetta, GA | August 2004 – February 2006"
    P[89].runs[0].text = (
        "Maintained and extended a legacy billing engine running on AIX and GNU/Linux using "
        "Oracle. The system was written primarily in C++ with shell scripts. Taught a C++ "
        "course to the code acceptance testing group and developed a Perl application that "
        "automated the configuration of their test boxes. Contract position."
    )
    P[89].runs[3].text = " C++, Perl, Shell, Oracle, AIX, GNU/Linux"

    # --- SAIC ---
    P[91].runs[0].text = "SAIC (Science Applications International Corporation)"
    P[92].runs[0].text = "Oak Ridge, TN | February 2003 – August 2004"
    replace_para_text(P[93], (
        "Maintained and extended a legacy Perl application for network-wide resource and "
        "account management at Oak Ridge and Los Alamos National Laboratories until it was "
        "decommissioned. Contract position."
    ))
    append_tech(P[93], "Perl, Oracle, Unix")

    # --- University of Tennessee ---
    P[96].runs[0].text = "Knoxville, TN | August 2002 – December 2002"
    P[96].runs[3].text = (
        "Built bioinformatics tools in Perl and C, backed by MySQL on Red Hat Linux, for "
        "monitoring, reporting, and analysis of microarray gene data. Tutored professional "
        "staff and graduate "
    )
    P[96].runs[4].text = "students in Perl programming for bioinformatics research. Contract position. "

    # --- Manhattan Associates ---
    P[99].runs[0].text = "Atlanta, GA | July 2001 – July 2002"
    replace_para_text(P[100], (
        "Performed code maintenance on a cross-platform warehouse management system written in "
        "C++ for Linux and Windows, using Oracle for data storage and the MKS Toolkit for "
        "Unix-style tooling on Windows. Contract position."
    ))
    append_tech(P[100], "C++, Oracle, Linux, Windows, MKS Toolkit")

    # --- LHS/Sema Group Telecom ---
    P[103].runs[0].text = "Atlanta, GA | January 2000 – July 2001"
    P[104].runs[0].text = (
        "Performed code maintenance and enhancement for a large legacy billing engine written "
        "in C/C++ using Oracle, spanning multiple Unix platforms. Contributed to major "
        "technical documentation rewrites. Contract position."
    )

    # --- Intellivoice / Premier Technologies ---
    P[107].runs[0].text = "Atlanta, GA | November 1998 – January 2000"
    P[108].runs[0].text = (
        "Worked on multimedia integration R&D for distributed Solaris networks. Developed "
        "coding projects using C++ and evaluated third-party middleware (such as CORBA) to "
        "facilitate network communication. Contract position."
    )

    # --- MCI/WorldCom ---
    P[111].runs[0].text = "Atlanta, GA | February 1996 – November 1998"
    P[112].runs[0].text = (
        "Full-cycle project to implement a billing engine on AIX using C++, Pro*C, and Oracle. "
        "Developed a Perl application that processed firearm purchase applications for the "
        "State of California (an MCI client). Contract position."
    )
    P[112].runs[3].text = " C++, Pro*C, Oracle, SQL, AIX, Unix, Perl"

    doc.save(str(DOCX_PATH))
    print(f"Saved updated document to {DOCX_PATH}")


if __name__ == "__main__":
    main()
