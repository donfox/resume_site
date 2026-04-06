"""
update_resume_docx.py

Updates Resume.v3.7.docx to match the current resume website text.
Run from the project root:

    conda run -n flask_env python scripts/update_resume_docx.py

A backup is written to static/files/Resume.v3.7.backup.docx before any changes.
"""

import shutil
from pathlib import Path
from docx import Document


DOCX_PATH = Path("static/files/Resume.v3.7.docx")
BACKUP_PATH = Path("static/files/Resume.v3.7.backup.docx")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def para_text(p) -> str:
    """Full text of a paragraph, concatenating all runs."""
    return "".join(r.text for r in p.runs)


def replace_para_text(p, new_text: str) -> None:
    """
    Replace all run text in a paragraph with new_text while
    preserving the paragraph style and the formatting of the first run.
    """
    runs = p.runs
    if not runs:
        p.add_run(new_text)
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""


def delete_paragraph(p) -> None:
    """Remove a paragraph from the document body."""
    p._element.getparent().remove(p._element)


def find_para(doc, snippet: str):
    """Return the first paragraph whose full text contains snippet."""
    for p in doc.paragraphs:
        if snippet in para_text(p):
            return p
    return None


# ---------------------------------------------------------------------------
# Individual update functions
# ---------------------------------------------------------------------------


def update_summary(doc):
    p5 = find_para(doc, "Software engineer with a background in telecommunications")
    p6 = find_para(doc, "I've mentored interns and worked on collaborative projects")
    new_text = (
        "Thirty years of professional software development spanning telecommunications, "
        "warehousing, healthcare IT, and data systems. Primary focus has been backend "
        "engineering on Unix/Linux platforms, with experience across a broad range of "
        "languages, databases, and system architectures. More recent work includes web "
        "development, functional programming, and applied AI/prompt engineering. "
        "Proof-of-concept development \u2014 particularly in healthcare IT \u2014 has been a "
        "recurring part of the work, as has occasional intern mentoring."
    )
    if p5:
        replace_para_text(p5, new_text)
    if p6:
        delete_paragraph(p6)


def update_emrts_intro(doc):
    p = find_para(doc, "Freelance developer for a healthcare IT startup")
    if p:
        replace_para_text(
            p,
            (
                "Software engineer working with a healthcare IT startup across six engagements, "
                "developing open-source tools for medical records, billing, provider enrollment, "
                "and claims. Work has centered on proof-of-concept development and R&D, often in "
                "collaboration with academic and professional partners. Contracts listed below fill "
                "the gaps between EMRTS engagements."
            ),
        )


def remove_emrts_old_note(doc):
    p = find_para(doc, "Note: Gaps between EMRTS engagements are covered")
    if p:
        delete_paragraph(p)


def update_emrts_interim_cspire(doc):
    p = find_para(doc, "Interim: Cspire contract")
    if p:
        replace_para_text(
            p,
            (
                "Interim contracts: CSpire, Jackson MS (Jul 2018\u2013Jan 2019); "
                "Orchid Orthopedic Solutions, Remote (Feb 2020\u2013Dec 2020); "
                "Quantuvos, Remote (Jul 2021\u2013Feb 2022) \u2014 see entries below."
            ),
        )


def update_emrts_dec2011_to_jan2012(doc):
    p = find_para(doc, "December 2011")
    if p:
        text = para_text(p)
        replace_para_text(p, text.replace("December 2011", "January 2012"))


def update_emrts_interim_hilton(doc):
    p = find_para(doc, "Interim: Hilton contract")
    if p:
        replace_para_text(
            p,
            (
                "Interim contracts: Hilton, Memphis TN (Apr 2012\u2013Aug 2012); "
                "Lowes Inc., Mooresville NC (Aug 2012\u2013Jun 2013) \u2014 see entries below."
            ),
        )


def update_emrts_interim_hp(doc):
    p = find_para(doc, "Interim: Hewlett-Packard contract")
    if p:
        replace_para_text(
            p,
            (
                "Interim contract: Hewlett-Packard, Nashville TN (May 2010\u2013Dec 2011) "
                "\u2014 see entry below."
            ),
        )


def update_quantuvos_dates(doc):
    p = find_para(doc, "Nov 2021")
    if p and "Remote" in para_text(p):
        text = para_text(p)
        replace_para_text(p, text.replace("Nov 2021", "Jul 2021"))


def update_quantuvos_description(doc):
    p = find_para(doc, "Reviewed and tested code for the MVP used by coaches")
    if p:
        replace_para_text(
            p,
            (
                "Technical consulting for Quantuvos, a coaching-platform startup. Reviewed and "
                "evaluated code for the client-facing MVP, supporting coaches and clients in "
                "scheduling workshops and sessions. Performed code reviews, identified issues, "
                "and recommended improvements. Contract position."
            ),
        )


def update_orchid_description(doc):
    p = find_para(doc, "Contract for document processing and archival")
    if p:
        replace_para_text(
            p,
            (
                "Contract work preparing Word, Excel, and PDF documents for archival submission "
                "to EtQ Reliance. Programmatically inserted tags and formatting in Word documents "
                "and validated spreadsheet metadata for accuracy and completeness. Contract position."
            ),
        )


def update_cspire_description(doc):
    p = find_para(doc, "RESTful auto-provisioning API for managing customer and vendor")
    if p:
        replace_para_text(
            p,
            (
                "Led development of a RESTful auto-provisioning API for tracking and managing "
                "customer services delivered through CSpire and its third-party vendor partners. "
                "The application supported logging, reporting, and administration for both business "
                "office and network teams. Contract position."
            ),
        )


def update_lowes_2012_description(doc):
    p = find_para(
        doc, "Refactored and extended legacy C code to support the transition to EMV"
    )
    if p:
        replace_para_text(
            p,
            (
                "Refactored and extended legacy C code to support EMV (chip-based credit and debit "
                "card) transaction processing as part of Lowe\u2019s transition away from magnetic "
                "stripe. Work included code analysis, bug fixes, and documentation. Contract position."
            ),
        )


def update_hilton_description(doc):
    p = find_para(
        doc, "Provided production support for Hilton's Central Reservation System"
    )
    if p:
        replace_para_text(
            p,
            (
                "Provided production support for Hilton\u2019s Central Reservation System, "
                "contributing bug fixes and maintenance using C, Perl, and shell scripting. "
                "Contract position."
            ),
        )


def update_acs_description(doc):
    p = find_para(
        doc,
        "Migrated MMIS (Medicaid Management Information System) software for TennCare from Hewlett",
    )
    if p:
        replace_para_text(
            p,
            (
                "Migrated MMIS (Medicaid Management Information System) software for TennCare from "
                "Hewlett-Packard\u2019s infrastructure to ACS. Developed automation scripts and "
                "analyzed legacy C modules. Contract position."
            ),
        )


def update_lexmark_description(doc):
    p = find_para(
        doc,
        "Performed bug fixes, code analysis, and C code refactoring for high-end prototype printers",
    )
    if p:
        replace_para_text(
            p,
            (
                "Bug fixes, code analysis, and refactoring of C code for high-end prototype printers "
                "designed for secure environments (government, military, finance). Developed Python "
                "scripts for daily integration testing. Contract position."
            ),
        )


def update_ibm_description(doc):
    p = find_para(
        doc,
        "Developed tools for porting and rebuilding C/C++ applications originally developed in Microsoft Visual Studio",
    )
    if p:
        replace_para_text(
            p,
            (
                "Developed tools for porting and rebuilding C/C++ applications from Microsoft Visual "
                "Studio to Unix environments, using Visual Studio build artifacts rather than "
                "traditional makefile-based builds. Contract position."
            ),
        )


def update_lowes_2006_description(doc):
    p = find_para(doc, "Internationalized software for Lowe's first Canadian store")
    if p:
        replace_para_text(
            p,
            (
                "Adapted Lowe\u2019s point-of-sale software for the company\u2019s first Canadian "
                "store, modifying C code to support Canadian currency, credit/debit card processing, "
                "and compliance with Canadian legal and regulatory requirements. Work done on AIX. "
                "Contract position."
            ),
        )


def update_saic_description(doc):
    p = find_para(
        doc,
        "Maintained and expanded the functionality of a legacy Perl application used for network-wide",
    )
    if p:
        replace_para_text(
            p,
            (
                "Maintained and extended a legacy Perl application for network-wide resource and "
                "account management at Oak Ridge and Los Alamos National Laboratories until it was "
                "decommissioned."
            ),
        )


def update_ut_dates(doc):
    p = find_para(doc, "Aug 2002 \u2013 Oct 2003")
    if p:
        for r in p.runs:
            if "Oct 2003" in r.text:
                r.text = r.text.replace("Oct 2003", "Dec 2002")
                break


def update_manhattan_description(doc):
    p = find_para(
        doc,
        "Performed code maintenance on a cross-platform warehouse management application",
    )
    if p:
        replace_para_text(
            p,
            (
                "Performed code maintenance on a cross-platform warehouse management system written "
                "in C++ for Linux and Windows. Contract position."
            ),
        )


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def main():
    if not DOCX_PATH.exists():
        raise FileNotFoundError(f"Not found: {DOCX_PATH}")

    shutil.copy2(DOCX_PATH, BACKUP_PATH)
    print(f"Backup written to {BACKUP_PATH}")

    doc = Document(str(DOCX_PATH))

    updates = [
        ("Summary", update_summary),
        ("EMRTS intro", update_emrts_intro),
        ("EMRTS old note removed", remove_emrts_old_note),
        ("EMRTS interim CSpire", update_emrts_interim_cspire),
        ("EMRTS Dec->Jan 2012", update_emrts_dec2011_to_jan2012),
        ("EMRTS interim Hilton", update_emrts_interim_hilton),
        ("EMRTS interim HP", update_emrts_interim_hp),
        ("Quantuvos dates", update_quantuvos_dates),
        ("Quantuvos description", update_quantuvos_description),
        ("Orchid description", update_orchid_description),
        ("CSpire description", update_cspire_description),
        ("Lowes 2012 description", update_lowes_2012_description),
        ("Hilton description", update_hilton_description),
        ("ACS description", update_acs_description),
        ("Lexmark description", update_lexmark_description),
        ("IBM ISS description", update_ibm_description),
        ("Lowe's 2006 description", update_lowes_2006_description),
        ("SAIC description", update_saic_description),
        ("UT Center dates", update_ut_dates),
        ("Manhattan description", update_manhattan_description),
    ]

    for label, fn in updates:
        try:
            fn(doc)
            print(f"  OK  {label}")
        except Exception as e:
            print(f"  FAIL {label}: {e}")

    doc.save(str(DOCX_PATH))
    print(f"\nSaved updated document to {DOCX_PATH}")


if __name__ == "__main__":
    main()
